"""Cuenta las filas de datos de un Excel exportado agrupadas por la columna "Fecha".

Solo imprime fechas y conteos (sin nombres, IDs ni horarios). Sirve para comparar con
los conteos por fecha de la consulta de la interfaz.

Uso: uv run --no-sync python exploracion/conteo_por_fecha.py <ruta.xlsx>
"""

import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def main(ruta: Path) -> None:
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = libro[libro.sheetnames[0]]
        idx_fecha = -1
        conteo: Counter[str] = Counter()
        filas_vacias_intermedias = 0
        en_datos = False
        for fila in hoja.iter_rows(values_only=True):
            textos = [str(c).strip() if c is not None else "" for c in fila]
            if not en_datos:
                if "Fecha" in textos and "Departamento" in textos:
                    idx_fecha = textos.index("Fecha")
                    en_datos = True
                continue
            if not any(textos):
                filas_vacias_intermedias += 1
                continue
            conteo[textos[idx_fecha] if idx_fecha < len(textos) else "?"] += 1
        print(
            f"filas de datos: {sum(conteo.values())}; filas vacías después de los encabezados: "
            f"{filas_vacias_intermedias}"
        )
        for fecha in sorted(conteo):
            print(f"  {fecha}: {conteo[fecha]}")
    finally:
        libro.close()


if __name__ == "__main__":
    main(Path(sys.argv[1]))
