"""Compara, fecha por fecha, la consulta de la interfaz con un Excel exportado.

En la sesión de exploración (iniciada por el usuario): va a "Registros de entradas y
salidas", elige el periodo y el departamento, recorre las páginas de la consulta y
cuenta registros por fecha. Luego cuenta las filas del Excel por fecha y compara.

Solo imprime fechas, números, nombres de campos y longitudes; nunca valores personales.
Los registros se comparan entre páginas con huellas SHA-256 en memoria.

Uso:
  uv run --no-sync python exploracion/comparar_por_fecha.py <ruta.xlsx> \
      [--periodo "Mes actual"] [--grupo Fisioterapeutas]
"""

import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from cdp import _pagina
from openpyxl import load_workbook
from playwright.sync_api import Page, Response, sync_playwright

API_REGISTROS = "/ISAPI/Bumblebee/AttendancePlugin/V1/Record"
REPORTE = "Registros de entradas y salidas"
FECHA = re.compile(r"^\d{4}-\d{2}-\d{2}")


def ir_al_reporte(page: Page) -> None:
    if page.locator("button.login-btn").filter(visible=True).count():
        raise SystemExit("No hay sesión: se ve la pantalla de login.")
    page.locator("[role='tab'].el-tabs__item").filter(
        has=page.get_by_text("Asistencia", exact=True)
    ).filter(visible=True).first.click()
    page.wait_for_timeout(1500)
    ok = page.locator("button.refresh-know").filter(visible=True)
    if ok.count():
        ok.first.click()
    menu = page.locator("ul.el-menu")
    item = menu.get_by_text(REPORTE, exact=True)
    if not item.filter(visible=True).count():
        menu.get_by_text("Registro de asistencia", exact=True).filter(visible=True).first.click()
    item.filter(visible=True).first.click()
    page.locator("button.export-btn").filter(visible=True).first.wait_for(timeout=30_000)
    if ok.count():
        ok.first.click()


def elegir_periodo(page: Page, periodo: str) -> None:
    candidatos = page.get_by_placeholder("Por favor, seleccione.").filter(visible=True)
    valores = [candidatos.nth(i).input_value() for i in range(candidatos.count())]
    indice = next(i for i, v in enumerate(valores) if v != "100")
    candidatos.nth(indice).click()
    page.locator(".el-select-dropdown").filter(visible=True).get_by_text(
        periodo, exact=True
    ).first.click()
    page.wait_for_timeout(1500)


def elegir_grupo(page: Page, departamento: str) -> Response:
    page.locator("button").filter(has_text=re.compile(r"personas seleccionadas\s*$", re.I)).filter(
        visible=True
    ).first.click()
    page.locator(f"li[class*='level'] a[title='{departamento}']").filter(visible=True).first.click()
    page.wait_for_timeout(1000)
    todos = (
        page.locator("label.el-checkbox")
        .filter(has_text="Seleccionar todos")
        .filter(visible=True)
        .first
    )
    if "is-checked" not in (todos.get_attribute("class") or ""):
        todos.click()
    with page.expect_response(lambda r: API_REGISTROS in r.url, timeout=30_000) as info:
        page.locator("button[title='Aceptar']").filter(visible=True).first.click()
    return info.value


def fecha_de(registro: dict[str, Any]) -> str:
    for v in registro.values():
        if isinstance(v, str) and FECHA.match(v):
            return v[:10]
    return "?"


def huella(registro: dict[str, Any]) -> str:
    return hashlib.sha256(json.dumps(registro, sort_keys=True).encode()).hexdigest()


def contar_excel(ruta: Path) -> Counter[str]:
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = libro[libro.sheetnames[0]]
        idx = -1
        conteo: Counter[str] = Counter()
        for fila in hoja.iter_rows(values_only=True):
            textos = [str(c).strip() if c is not None else "" for c in fila]
            if idx < 0:
                if "Fecha" in textos and "Departamento" in textos:
                    idx = textos.index("Fecha")
                continue
            if any(textos):
                conteo[textos[idx]] += 1
        return conteo
    finally:
        libro.close()


def main() -> None:
    a = argparse.ArgumentParser()
    a.add_argument("excel", type=Path)
    a.add_argument("--periodo", default="Mes actual")
    a.add_argument("--grupo", default="Fisioterapeutas")
    args = a.parse_args()

    paginas: list[dict[str, Any]] = []
    registros: list[dict[str, Any]] = []
    with sync_playwright() as p:
        page = _pagina(p)
        ir_al_reporte(page)
        elegir_periodo(page, args.periodo)
        respuesta = elegir_grupo(page, args.grupo)
        while True:
            consulta = json.loads(respuesta.request.post_data or "{}")["RecordRequest"]
            datos = respuesta.json()["ResponseStatus"]["Data"]
            lote = datos.get("OnOffDutyRecord") or []
            paginas.append(
                {
                    "pedida": consulta.get("PageIndex"),
                    "tamano_pedido": consulta.get("PageSize"),
                    "devuelta": datos.get("PageIndex"),
                    "registros": len(lote),
                    "NextPage": datos.get("NextPage"),
                    "rango": f"{consulta['QueryInfo']['BeginTime']} → {consulta['QueryInfo']['EndTime']}",
                    "personas": len(consulta["QueryInfo"]["PersonID"]),
                }
            )
            registros += lote
            if datos.get("NextPage") != 1:
                break
            with page.expect_response(lambda r: API_REGISTROS in r.url, timeout=30_000) as info:
                page.locator("button.btn-next").filter(visible=True).first.click()
            respuesta = info.value

    print("páginas de la consulta:")
    for pg in paginas:
        print(f"  {json.dumps(pg, ensure_ascii=False)}")
    huellas = Counter(huella(r) for r in registros)
    print(
        f"registros sumados: {len(registros)}; distintos: {len(huellas)}; "
        f"repetidos entre páginas: {sum(n - 1 for n in huellas.values() if n > 1)}"
    )
    if registros:
        tipos = {k: type(v).__name__ for k, v in registros[0].items()}
        print(f"campos del registro: {json.dumps(tipos, ensure_ascii=False)}")

    api = Counter(fecha_de(r) for r in registros)
    excel = contar_excel(args.excel)
    print(f"\n{'fecha':<12}{'interfaz':>9}{'excel':>7}{'dif':>6}")
    for f in sorted(set(api) | set(excel)):
        dif = excel.get(f, 0) - api.get(f, 0)
        marca = "  <--" if dif else ""
        print(f"{f:<12}{api.get(f, 0):>9}{excel.get(f, 0):>7}{dif:>6}{marca}")
    print(f"{'total':<12}{sum(api.values()):>9}{sum(excel.values()):>7}")


if __name__ == "__main__":
    main()
