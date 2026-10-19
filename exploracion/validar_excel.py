"""Valida la estructura de un Excel exportado sin imprimir datos personales.

Solo extrae: hojas, texto de "Periodo:", encabezados, conjunto de departamentos,
conteo de filas de datos, tamaño y SHA-256. Nunca nombres, apellidos, IDs ni horarios.

Uso: uv run --frozen python exploracion/validar_excel.py <ruta.xlsx>
"""

import hashlib
import sys
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

ENCABEZADOS_CLAVE = {"Nombre", "Departamento"}


@dataclass
class Estructura:
    archivo: str
    tamano: int
    sha256: str
    hojas: list[str]
    fila_periodo: int | None = None
    periodo: str = ""
    fila_encabezados: int | None = None
    encabezados: list[str] = field(default_factory=list)
    departamentos: set[str] = field(default_factory=set)
    filas_datos: int = 0
    max_fila: int | None = None


def inspeccionar(ruta: Path) -> Estructura:
    datos = ruta.read_bytes()
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = libro[libro.sheetnames[0]]
        e = Estructura(ruta.name, len(datos), hashlib.sha256(datos).hexdigest(), libro.sheetnames)
        idx_depto = -1
        for numero, fila in enumerate(hoja.iter_rows(values_only=True), start=1):
            textos = [str(c).strip() if c is not None else "" for c in fila]
            if e.fila_encabezados is None:
                for t in textos:
                    if t.startswith("Periodo:"):
                        e.fila_periodo, e.periodo = numero, t
                if ENCABEZADOS_CLAVE.issubset(set(textos)):
                    e.fila_encabezados = numero
                    e.encabezados = [t for t in textos if t]
                    idx_depto = textos.index("Departamento")
                continue
            if not any(textos):
                break
            e.filas_datos += 1
            if 0 <= idx_depto < len(textos):
                e.departamentos.add(textos[idx_depto])
        e.max_fila = hoja.max_row
        return e
    finally:
        libro.close()


def main(ruta: Path) -> None:
    e = inspeccionar(ruta)
    print(f"archivo: {e.archivo}  tamaño: {e.tamano} B  sha256: {e.sha256}")
    print(f"hojas: {e.hojas}")
    print(f"fila {e.fila_periodo}: {e.periodo!r}")
    print(f"fila {e.fila_encabezados}: encabezados = {e.encabezados}")
    print(f"departamentos = {sorted(e.departamentos)}")
    print(f"filas de datos = {e.filas_datos}  (dimensión declarada: {e.max_fila} filas)")


if __name__ == "__main__":
    main(Path(sys.argv[1]))
