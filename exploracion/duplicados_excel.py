"""Busca filas duplicadas y repeticiones persona+fecha en un Excel exportado.

Compara huellas (SHA-256) de filas completas y de la clave (ID, Fecha) sin imprimirlas.
Solo imprime conteos, fechas y los valores de "Periodo de tiempo" de las repeticiones.

Uso: uv run --no-sync python exploracion/duplicados_excel.py <ruta.xlsx>
"""

import hashlib
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


def huella(*partes: str) -> str:
    return hashlib.sha256("\x1f".join(partes).encode()).hexdigest()


def main(ruta: Path) -> None:
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = libro[libro.sheetnames[0]]
        cab: list[str] = []
        filas: list[list[str]] = []
        for fila in hoja.iter_rows(values_only=True):
            textos = [str(c).strip() if c is not None else "" for c in fila]
            if not cab:
                if "Fecha" in textos and "Departamento" in textos:
                    cab = textos
                continue
            if any(textos):
                filas.append(textos)
    finally:
        libro.close()

    i_id, i_fecha, i_periodo = cab.index("ID"), cab.index("Fecha"), cab.index("Periodo de tiempo")
    completas = Counter(huella(*f) for f in filas)
    exactas = sum(n - 1 for n in completas.values() if n > 1)
    print(f"filas: {len(filas)}; filas idénticas repetidas (sobrantes): {exactas}")

    por_clave: dict[str, list[list[str]]] = defaultdict(list)
    for f in filas:
        por_clave[huella(f[i_id], f[i_fecha])].append(f)
    repetidas = {k: v for k, v in por_clave.items() if len(v) > 1}
    sobrantes = sum(len(v) - 1 for v in repetidas.values())
    print(
        f"claves persona+fecha distintas: {len(por_clave)}; claves con más de una fila: "
        f"{len(repetidas)}; filas sobrantes por esa causa: {sobrantes}"
    )
    for grupo in repetidas.values():
        identicas = len({huella(*f) for f in grupo}) == 1
        periodos = [f[i_periodo] for f in grupo]
        print(
            f"  fecha {grupo[0][i_fecha]}: {len(grupo)} filas; idénticas={identicas}; "
            f"'Periodo de tiempo'={periodos}"
        )


if __name__ == "__main__":
    main(Path(sys.argv[1]))
