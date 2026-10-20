"""Lee la estructura del Excel exportado con openpyxl en solo lectura.

No depende de números de fila: busca la fila "Periodo:" y la de encabezados. Las claves
(ID, Fecha) quedan solo en memoria para la validación de ADR-0005; no se registran.
"""

import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from asistencia_bot.domain.errores import ErrorLectura
from asistencia_bot.domain.modelos import REPORTE_HIK, ClaveRegistro, MetadatosReporte

PREFIJO_PERIODO = "Periodo:"
ENCABEZADOS_ANCLA = ("Nombre", "Departamento", "Fecha")
FILAS_MAXIMAS_ENCABEZADO = 40


def _texto(valor: object) -> str:
    return str(valor).strip() if valor is not None else ""


def _fecha(valor: object) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    try:
        return date.fromisoformat(texto[:10]) if texto else None
    except ValueError:
        return None


class InspectorExcel:
    def leer_metadatos(self, ruta: Path) -> MetadatosReporte:
        try:
            libro = load_workbook(ruta, read_only=True, data_only=True)
        except (OSError, zipfile.BadZipFile, InvalidFileException, KeyError) as e:
            raise ErrorLectura(f"no se pudo abrir {ruta.name}: {type(e).__name__}") from None
        try:
            hoja = libro[libro.sheetnames[0]]
            titulo = periodo = ""
            encabezados: list[str] = []
            departamentos: set[str] = set()
            claves: set[ClaveRegistro] = set()
            filas = sin_fecha = 0
            i_id = i_depto = i_fecha = -1
            for numero, fila in enumerate(hoja.iter_rows(values_only=True), start=1):
                textos = [_texto(c) for c in fila]
                if not encabezados:
                    for t in textos:
                        if t.startswith(PREFIJO_PERIODO):
                            periodo = t.removeprefix(PREFIJO_PERIODO).strip()
                        elif t.startswith(REPORTE_HIK) and not titulo:
                            titulo = t
                    if all(a in textos for a in ENCABEZADOS_ANCLA):
                        encabezados = textos
                        i_depto, i_fecha = textos.index("Departamento"), textos.index("Fecha")
                        i_id = textos.index("ID") if "ID" in textos else -1
                    elif numero >= FILAS_MAXIMAS_ENCABEZADO:
                        break
                    continue
                if not any(textos):
                    break
                filas += 1
                if i_depto < len(textos):
                    departamentos.add(textos[i_depto])
                dia = _fecha(fila[i_fecha]) if i_fecha < len(fila) else None
                if dia is None:
                    sin_fecha += 1
                elif 0 <= i_id < len(textos):
                    claves.add((textos[i_id], dia))
            return MetadatosReporte(
                hoja=libro.sheetnames[0],
                titulo=titulo,
                periodo_texto=periodo,
                encabezados=tuple(e for e in encabezados if e),
                departamentos=frozenset(departamentos),
                filas=filas,
                claves=frozenset(claves),
                filas_sin_fecha=sin_fecha,
            )
        finally:
            libro.close()
